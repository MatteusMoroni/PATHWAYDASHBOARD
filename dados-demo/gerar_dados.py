# Gera os dados FICTÍCIOS da demo PC 102 - Block 4 (5 exports .xlsx + backup .json)
# e valida tudo relendo os arquivos e simulando o app no dia D.
# Uso: uv run --no-project --with openpyxl python -I gerar_dados.py
import json, os, random, sys, unicodedata
from datetime import date
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

OUT = os.path.dirname(os.path.abspath(__file__))
TURMA = "PC 102 - Block 4"
COLS = ["Full Name", "Student Number", "Gathering", "Gender", "Email (Contact)", "Email Address 2",
        "Mobile Phone", "Age", "Last Participation", "City", "Country", "Is Member"]
EXPORTS = ["2026-09-24", "2026-10-01", "2026-10-05", "2026-10-08", "2026-10-14"]
GATHER = ["2026-09-09", "2026-09-16", "2026-09-23", "2026-09-30", "2026-10-07", "2026-10-14"]
D = date(2026, 10, 15)
rnd = random.Random(1015)

def ascii_(s):
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()

# S, nome, gênero, idade, cidade, DDD, membro, tem email2
CAST = [
    (1, "Ana Clara Moreira", "Female", 24, "São Paulo", "11", "Yes", True),
    (2, "Bruno Teixeira Lima", "Male", 31, "Campinas", "19", "Yes", True),
    (3, "Camila Duarte Ribeiro", "Female", 27, "Rio de Janeiro", "21", "No", False),
    (4, "Daniel Fonseca Prado", "Male", 45, "Belo Horizonte", "31", "Yes", True),
    (5, "Eduarda Martins Rocha", "Female", 22, "Curitiba", "41", "Yes", True),
    (6, "Felipe Nogueira Batista", "Male", 38, "Porto Alegre", "51", "Yes", True),
    (7, "Gabriela Souza Pires", "Female", 29, "Florianópolis", "48", "Yes", False),
    (8, "Henrique Campos Vidal", "Male", 52, "Salvador", "71", "No", True),
    (9, "Isabela Cardoso Lemos", "Female", 33, "Recife", "81", "Yes", True),
    (10, "Juliano Barros Machado", "Male", 26, "Fortaleza", "85", "Yes", True),
    (11, "Larissa Mendes Coelho", "Female", 41, "Goiânia", "62", "Yes", True),
    (12, "Marcos Vieira Antunes", "Male", 19, "Brasília", "61", "Yes", True),
    (13, "Natália Freitas Rangel", "Female", 35, "Manaus", "92", "No", True),
    (14, "Otávio Brandão Siqueira", "Male", 28, "Belém", "91", "Yes", True),  # email2 só a partir de E4
    (15, "Patrícia Lopes Andrade", "Female", 47, "Natal", "84", "Yes", True),
    (16, "Rodrigo Pacheco Dias", "Male", 23, "Joinville", "47", "No", True),
    (17, "Sabrina Monteiro Reis", "Female", 30, "Londrina", "43", "Yes", True),
    (18, "Thiago Albuquerque Neves", "Male", 20, "Vitória", "27", "Yes", False),
    (19, "Vanessa Correia Tavares", "Female", 39, "Uberlândia", "34", "Yes", True),
    (20, "Wagner Matos Figueira", "Male", 44, "Santos", "13", "No", True),
    (21, "Yasmin Carvalho Peixoto", "Female", 21, "Ribeirão Preto", "16", "Yes", False),
    (22, "Leandro Quintela Farias", "Male", 34, "Maceió", "82", "No", True),
    (23, "Kátia Sampaio Furtado", "Female", 25, "João Pessoa", "83", "Yes", True),
]

def mat(n): return f"990000{n:02d}"

def base(n):
    _, nome, g, idade, cid, ddd, mem, e2 = CAST[n - 1]
    p = ascii_(nome).split()
    return {"nome": nome, "matricula": mat(n), "genero": g, "idade": idade, "cidade": cid, "ddd": ddd,
            "membro": mem, "email": f"{p[0]}.{p[-1]}@pathway-demo.example",
            "email2": f"{p[0]}{p[-1]}{idade}@mail-demo.example" if e2 else ""}

def cadastro(n, ei):  # ei = índice do export (0..4)
    b = base(n)
    tel = f"55{b['ddd']}99999{n:04d}"
    if n == 9 and ei >= 2: tel = f"55{b['ddd']}99998{n:04d}"   # S09 troca de número a partir de E3
    if n == 14: b["email2"] = "otavio.siqueira.pessoal@mail-demo.example" if ei >= 3 else ""
    b["telefone"] = tel
    return b

def d(s):  # "dd/mm" -> ISO
    dd, mm = s.split("/"); return f"2026-{mm}-{dd}"

LP = {  # E1..E5; None = vazio; "X" = aluno ausente do export
    1: "24/09 01/10 05/10 08/10 14/10", 2: "23/09 30/09 04/10 07/10 13/10",
    3: "22/09 29/09 03/10 06/10 13/10", 4: "24/09 30/09 05/10 08/10 14/10",
    5: "23/09 01/10 04/10 08/10 13/10", 6: "24/09 29/09 05/10 07/10 14/10",
    7: "22/09 01/10 03/10 07/10 14/10", 8: "23/09 30/09 04/10 06/10 13/10",
    9: "24/09 01/10 05/10 08/10 13/10",
    10: "23/09 27/09 29/09 07/10 13/10", 11: "21/09 26/09 28/09 28/09 28/09",
    12: "23/09 24/09 03/10 07/10 13/10", 13: "23/09 27/09 27/09 07/10 13/10",
    14: "22/09 28/09 28/09 28/09 12/10", 15: "23/09 01/10 01/10 01/10 12/10",
    16: "21/09 25/09 30/09 04/10 04/10", 17: "23/09 30/09 03/10 07/10 10/10",
    18: "- - - - -", 19: "22/09 30/09 04/10 X 13/10",
    20: "07/09 07/09 07/09 07/09 07/09", 21: "15/09 15/09 15/09 15/09 15/09",
    22: "20/09 24/09 26/09 X X", 23: "X X 04/10 07/10 13/10",
}
LP = {k: [None if t == "-" else ("X" if t == "X" else d(t)) for t in v.split()] for k, v in LP.items()}

def present(n, ei): return LP[n][ei] != "X"

CHAMADA = {n: "PPPPPP" for n in range(1, 9)}
CHAMADA[5] = "PPJPPP"; CHAMADA[7] = "PPPPJP"
CHAMADA.update({9: "PPPPPF", 10: "PPFPPP", 11: "PPPPFF", 12: "PPPFPP", 13: "PPPFFP", 14: "PPPPFP",
                15: "PPPPFP", 16: "PPPPFF", 17: "PPPPPF", 18: "PPPPPF", 19: "PPPPPP", 20: "PFFFFF",
                21: "PPFFFF", 22: "PPPFFF", 23: "....PP"})

CONTATOS = {
    2: [("12/10", "Presencial", "Falou com o aluno", "Conversamos depois da reunião da ala; está animado com a turma.")],
    6: [("13/10", "WhatsApp", "Falou com o aluno", "Pediu dica de como organizar as tarefas da semana.")],
    10: [("05/10", "WhatsApp", "Falou com o aluno", "Faltou dia 23/09 por plantão no trabalho; vai repor a leitura.")],
    11: [("02/10", "WhatsApp", "Não atendeu", "Mensagem enviada, sem visualização."),
         ("06/10", "Ligação", "Deixou recado", "Caixa postal; deixamos recado pedindo retorno."),
         ("08/10", "WhatsApp", "Prometeu voltar", "Respondeu que a mãe está internada; promete voltar após a alta.")],
    12: [("02/10", "Ligação", "Falou com o aluno", "Não sabia enviar a tarefa no Canvas; mostramos passo a passo."),
         ("12/10", "WhatsApp", "Falou com o aluno", "Confirmou que entregou tudo da semana.")],
    13: [("06/10", "WhatsApp", "Não atendeu", "Mensagem sem resposta."),
         ("07/10", "Ligação", "Falou com o aluno", "Ficou sem internet em casa; vai usar o wi-fi da capela."),
         ("13/10", "Mensagem no Canvas", "Falou com o aluno", "Internet voltou; já está em dia com as tarefas.")],
    14: [("06/10", "E-mail", "Deixou recado", "E-mail perguntando se está tudo bem."),
         ("09/10", "WhatsApp", "Falou com o aluno", "Mudou de emprego e perdeu o ritmo; combinamos um horário fixo de estudo."),
         ("13/10", "WhatsApp", "Falou com o aluno", "Mandou novo e-mail pessoal; voltou a acessar o Canvas.")],
    15: [("09/10", "Ligação", "Falou com o aluno", "Achou que tinha perdido o prazo e desanimou; explicamos a política de atraso."),
         ("14/10", "WhatsApp", "Falou com o aluno", "Entregou as tarefas atrasadas e agradeceu.")],
    20: [("17/09", "Ligação", "Vai sair do curso", "Vai priorizar o trabalho neste semestre; pensa em voltar no próximo bloco.")],
    21: [("25/09", "WhatsApp", "Vai sair do curso", "Mudança de cidade; vai trancar e tentar no próximo ano.")],
    22: [("30/09", "WhatsApp", "Não atendeu", "Mensagem não entregue."),
         ("03/10", "Ligação", "Não atendeu", "Telefone chama e cai."),
         ("07/10", "Outro", "Deixou recado", "Pedimos ao líder da ala para avisar que estamos procurando por ele.")],
}
STATUS = {11: "atencao", 16: "risco", 17: "atencao", 22: "risco", 20: "desistente", 21: "desistente"}
NOTAS = {
    11: "Mãe internada. Ligar de novo depois da alta.",
    12: "Tinha dificuldade com o Canvas. Já resolvido.",
    13: "Problema de internet. Pode usar o wi-fi da capela.",
    14: "Novo emprego. Combinamos estudar de terça e quinta à noite.",
    16: "Parou de acessar em 04/10 e faltou duas vezes. Ainda não ligamos.",
    17: "Faltou ao último encontro. Mandar mensagem nesta semana.",
    20: "Saiu para priorizar o trabalho. Convidar para o próximo bloco.",
    21: "Mudou de cidade. Pode voltar no ano que vem.",
    22: "Sumiu do export em 08/10. Pedir ajuda ao líder da ala.",
}
TEMAS = ["Boas-vindas e metas do bloco", "Gestão do tempo", "Hábitos de estudo",
         "Comunicação eficaz", "Autossuficiência e carreira", "Revisão e próximos passos"]

def rid(): return "".join(rnd.choice("abcdefghijklmnopqrstuvwxyz0123456789") for _ in range(8))

# ---------- 1. Gera os .xlsx ----------
def sortkey(nome): return ascii_(nome)

files = []
for ei, em in enumerate(EXPORTS):
    wb = Workbook(); ws = wb.active; ws.title = "Student List"
    ws.append(COLS)
    for c in ws[1]: c.font = Font(bold=True)
    rows = []
    for n in range(1, 24):
        if not present(n, ei): continue
        c = cadastro(n, ei)
        rows.append([c["nome"], c["matricula"], TURMA, c["genero"], c["email"], c["email2"] or None,
                     c["telefone"], c["idade"], LP[n][ei], c["cidade"], "Brazil", c["membro"]])
    rows.sort(key=lambda r: sortkey(r[0]))
    for r in rows: ws.append(r)
    for col, w in zip("ABCDEFGHIJKL", [28, 15, 18, 9, 36, 40, 16, 6, 17, 16, 9, 10]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.column != 8: c.number_format = "@"  # tudo texto, exceto Age
    ws.freeze_panes = "A2"
    fn = os.path.join(OUT, f"export_{ei + 1:02d}_{em}.xlsx"); wb.save(fn); files.append(fn)

# ---------- 2. Simula o app importando os ARQUIVOS (não os dicionários) ----------
erros = []
def check(cond, msg):
    if not cond: erros.append(msg)

alunos, importacoes = {}, []
prev_lp = {}
for ei, fn in enumerate(files):
    wb = load_workbook(fn)
    check(wb.sheetnames == ["Student List"], f"{fn}: abas {wb.sheetnames}")
    ws = wb["Student List"]
    hdr = [c.value for c in ws[1]]
    check(hdr == COLS, f"{fn}: cabeçalho {hdr}")
    data = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    check([sortkey(r[0]) for r in data] == sorted(sortkey(r[0]) for r in data), f"{fn}: ordem alfabética")
    em = EXPORTS[ei]; novos = atual = 0; vistos = set()
    for r in data:
        rec = dict(zip(COLS, r))
        m = rec["Student Number"]
        check(isinstance(m, str) and len(m) == 8 and m.isdigit() and m.startswith("990000"), f"matrícula {m}")
        check(rec["Gathering"] == TURMA and rec["Country"] == "Brazil", f"{m}: turma/país")
        check(rec["Gender"] in ("Male", "Female") and rec["Is Member"] in ("Yes", "No"), f"{m}: gênero/membro")
        check(isinstance(rec["Age"], int) and 19 <= rec["Age"] <= 52, f"{m}: idade")
        tel = rec["Mobile Phone"]
        check(isinstance(tel, str) and tel.isdigit() and len(tel) == 13 and tel[4] == "9", f"{m}: telefone {tel}")
        lp = rec["Last Participation"]
        check(lp is None or (isinstance(lp, str) and len(lp) == 10 and lp <= em), f"{m}: LP {lp!r} em {em}")
        check(lp is None or prev_lp.get(m) is None or lp >= prev_lp[m], f"{m}: LP diminuiu")
        if lp: prev_lp[m] = lp
        check(rec["Email (Contact)"].endswith("@pathway-demo.example"), f"{m}: email")
        check(rec["Email Address 2"] in (None,) or rec["Email Address 2"].endswith("@mail-demo.example"), f"{m}: email2")
        vistos.add(m)
        a = alunos.get(m)
        if a is None:
            novos += 1
            a = alunos[m] = {"id": m, "contatos": [], "historico": [], "status": "ativo", "notas": "",
                             "saiuDoExport": False, "entrouEm": em, "ultimoAcesso": None}
        else:
            atual += 1
            for k, col in [("nome", "Full Name"), ("genero", "Gender"), ("idade", "Age"),
                           ("cidade", "City"), ("membro", "Is Member")]:
                check(a[k] == rec[col], f"{m}: {k} mudou")
        a.update({"nome": rec["Full Name"], "matricula": m, "email": rec["Email (Contact)"],
                  "email2": rec["Email Address 2"] or "", "telefone": tel, "idade": rec["Age"],
                  "genero": rec["Gender"], "cidade": rec["City"], "pais": "Brazil",
                  "membro": rec["Is Member"], "turma": TURMA})
        if lp:
            a["ultimoAcesso"] = lp; a["historico"].append({"em": em, "acesso": lp})
        if a["saiuDoExport"]:
            a["saiuDoExport"] = False  # o app mantém saiuEm (Dashboard.html, importar)
    sumiram = 0
    for m, a in alunos.items():
        if m not in vistos and not a["saiuDoExport"]:
            a["saiuDoExport"] = True; a["saiuEm"] = em; sumiram += 1
    importacoes.insert(0, {"em": em, "total": len(data), "novos": novos, "atualizados": atual, "sumiram": sumiram})

# DDD coerente e unicidade
nomes = [a["nome"] for a in alunos.values()]; emails = [a["email"] for a in alunos.values()]
check(len(set(nomes)) == 23 and len(set(emails)) == 23, "nomes/e-mails duplicados")
for n in range(1, 24):
    check(alunos[mat(n)]["telefone"][2:4] == CAST[n - 1][5], f"S{n:02d}: DDD")
membros = sum(a["membro"] == "Yes" for a in alunos.values())

# registra chamada, contatos, status e notas como o casal faria no app
for n in range(1, 24):
    a = alunos[mat(n)]
    a["status"] = STATUS.get(n, "ativo"); a["notas"] = NOTAS.get(n, "")
    a["contatos"] = [{"id": rid(), "data": d(dt), "canal": c, "resultado": r, "obs": o}
                     for dt, c, r, o in CONTATOS.get(n, [])]
encontros = []
for gi, g in enumerate(GATHER):
    pres = {mat(n): CHAMADA[n][gi] for n in range(1, 24) if CHAMADA[n][gi] != "."}
    encontros.append({"id": rid(), "data": g, "tema": TEMAS[gi], "presencas": pres, "aberto": True})

ORDEM = ["id", "nome", "matricula", "email", "email2", "telefone", "idade", "genero", "cidade", "pais",
         "membro", "turma", "ultimoAcesso", "historico", "status", "notas", "contatos", "saiuDoExport",
         "saiuEm", "entrouEm"]
estado = {"versao": 1, "turma": TURMA, "cfg": {"canvas": 7, "contato": 7},
          "alunos": {m: {k: alunos[m][k] for k in ORDEM if k in alunos[m]} for m in sorted(alunos)},
          "encontros": encontros, "importacoes": importacoes}
jfn = os.path.join(OUT, "pc102-demo-backup.json")
with open(jfn, "w", encoding="utf-8") as f: json.dump(estado, f, ensure_ascii=False, indent=2)

# ---------- 3. Consistência da história ----------
for n in range(1, 24):
    a = estado["alunos"][mat(n)]
    for ct in a["contatos"]:
        check(ct["data"] <= "2026-10-14", f"S{n:02d}: contato no futuro")
# retorno ao Canvas depois do primeiro contato (casos recuperados)
for n in (10, 12, 13, 14, 15):
    a = estado["alunos"][mat(n)]; c0 = a["contatos"][0]["data"]
    check(a["ultimoAcesso"] > c0, f"S{n:02d}: não voltou depois do contato")

# ---------- 4. Painel em D (relê o JSON) ----------
E = json.load(open(jfn, encoding="utf-8"))
ult_enc = sorted(E["encontros"], key=lambda e: e["data"])
def dias(iso): return (D - date.fromisoformat(iso)).days
def faltas_seguidas(m):
    k = 0
    for e in reversed(ult_enc):
        if e["presencas"].get(m) == "F": k += 1
        else: break
    return k
painel = []
for m, a in E["alunos"].items():
    if a["status"] == "desistente": continue
    dc = dias(a["ultimoAcesso"]) if a["ultimoAcesso"] else None
    fs = faltas_seguidas(m)
    dct = dias(max(c["data"] for c in a["contatos"])) if a["contatos"] else None
    p = 0
    if dc is not None and dc >= 7: p += 3
    elif dc is not None and dc >= 4: p += 1
    p += 3 if fs >= 2 else (1 if fs == 1 else 0)
    if dct is None and ((dc or 0) >= 7 or fs >= 1): p += 2
    elif dct is not None and dct >= 7: p += 1
    if a["status"] == "risco": p += 2
    if a["saiuDoExport"]: p += 3
    painel.append((m, p, dc, fs, dct, a))
lista = sorted([x for x in painel if x[1] > 0],
               key=lambda x: (-x[1], -(x[2] if x[2] is not None else -1), -x[3], -(x[4] if x[4] is not None else 9999)))
S = lambda m: "S" + m[-2:]
ind = {
    "sem_canvas_7": sorted(S(m) for m, p, dc, fs, dct, a in painel if dc is not None and dc >= 7),
    "faltaram_ultimo": sorted(S(m) for m, p, dc, fs, dct, a in painel if fs >= 1),
    "ativos": len(painel),
    "sem_contato_7": sum(1 for x in painel if x[4] is None or x[4] >= 7),
}

esperado_imp = [("2026-09-24", 22, 22, 0, 0), ("2026-10-01", 22, 0, 22, 0), ("2026-10-05", 23, 1, 22, 0),
                ("2026-10-08", 21, 0, 21, 2), ("2026-10-14", 22, 0, 22, 0)]
obtido_imp = [(i["em"], i["total"], i["novos"], i["atualizados"], i["sumiram"]) for i in reversed(E["importacoes"])]
check(obtido_imp == esperado_imp, f"importações {obtido_imp}")
check(ind["sem_canvas_7"] == ["S11", "S16", "S22"], f"sem canvas {ind['sem_canvas_7']}")
check(ind["faltaram_ultimo"] == ["S09", "S11", "S16", "S17", "S18", "S22"], f"faltaram {ind['faltaram_ultimo']}")
check(ind["ativos"] == 21, f"ativos {ind['ativos']}")
check(ind["sem_contato_7"] == 15, f"sem contato {ind['sem_contato_7']}")
check([(S(x[0]), x[1]) for x in lista] == [("S22", 12), ("S16", 10), ("S11", 7), ("S17", 4), ("S09", 3), ("S18", 3), ("S10", 1)],
      f"lista {[(S(x[0]), x[1]) for x in lista]}")
canais = {c["canal"] for a in E["alunos"].values() for c in a["contatos"]}
resultados = {c["resultado"] for a in E["alunos"].values() for c in a["contatos"]}
check(len(canais) == 6 and len(resultados) == 5, "canais/resultados incompletos")
check(E["alunos"]["99000018"]["ultimoAcesso"] is None and not E["alunos"]["99000018"]["historico"], "S18")
check(E["alunos"]["99000022"].get("saiuEm") == "2026-10-08" and E["alunos"]["99000022"]["saiuDoExport"], "S22 saiu")
check(sum(a["saiuDoExport"] for a in E["alunos"].values()) == 1, "saiuDoExport além de S22")
check(E["alunos"]["99000019"].get("saiuEm") == "2026-10-08" and not E["alunos"]["99000019"]["saiuDoExport"], "S19 voltou")

# ---------- relatório ----------
print("Importações:", obtido_imp)
print("Indicadores:", ind, "| membros Yes:", membros, "de 23")
for i, (m, p, dc, fs, dct, a) in enumerate(lista, 1):
    print(f"{i}. {S(m)} {a['nome']}: {p} pts | canvas={dc} faltas={fs} contato={dct} fora={a['saiuDoExport']}")
print("Canais:", sorted(canais)); print("Resultados:", sorted(resultados))
print("ERROS:" if erros else "OK: todas as validações passaram")
for e in erros: print(" -", e)
sys.exit(1 if erros else 0)
