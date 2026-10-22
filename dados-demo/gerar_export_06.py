# Gera export_06_2026-10-15.xlsx: a lista do E5 (idêntica) + 3 alunos novos FICTÍCIOS (S24 a S26).
# Serve para mostrar o alerta "aluno novo" ao importar no dia da demo.
# Uso: uv run --no-project --with openpyxl python -I gerar_export_06.py
import os, sys, unicodedata
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

OUT = os.path.dirname(os.path.abspath(__file__))
E5 = os.path.join(OUT, "export_05_2026-10-14.xlsx")
E6 = os.path.join(OUT, "export_06_2026-10-15.xlsx")
TURMA = "PC 102 - Block 4"

def ascii_(s):
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()

# S, nome, gênero, idade, cidade, DDD, membro, tem email2, Last Participation
NOVOS = [
    (24, "Mariana Toledo Bastos", "Female", 23, "Aracaju", "79", "Yes", True, "2026-10-15"),
    (25, "Gustavo Ferraz Lacerda", "Male", 37, "Cuiabá", "65", "Yes", True, "2026-10-14"),
    (26, "Priscila Nunes Gouveia", "Female", 31, "Teresina", "86", "No", False, None),  # ainda não entrou no Canvas
]

src = load_workbook(E5)["Student List"]
cols = [c.value for c in src[1]]
base = [[c.value for c in r] for r in src.iter_rows(min_row=2)]

novas = []
for n, nome, g, idade, cid, ddd, mem, e2, lp in NOVOS:
    p = ascii_(nome).split()
    novas.append([nome, f"990000{n:02d}", TURMA, g, f"{p[0]}.{p[-1]}@pathway-demo.example",
                  f"{p[0]}{p[-1]}{idade}@mail-demo.example" if e2 else None,
                  f"55{ddd}99999{n:04d}", idade, lp, cid, "Brazil", mem])

rows = sorted(base + novas, key=lambda r: ascii_(r[0]))
wb = Workbook(); ws = wb.active; ws.title = "Student List"
ws.append(cols)
for c in ws[1]: c.font = Font(bold=True)
for r in rows: ws.append(r)
for col, w in zip("ABCDEFGHIJKL", [28, 15, 18, 9, 36, 40, 16, 6, 17, 16, 9, 10]):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        if c.column != 8: c.number_format = "@"
ws.freeze_panes = "A2"
wb.save(E6)

# ---------- validação relendo o arquivo ----------
erros = []
def check(ok, msg):
    if not ok: erros.append(msg)

wb2 = load_workbook(E6); ws2 = wb2["Student List"]
check(wb2.sheetnames == ["Student List"], "abas")
check([c.value for c in ws2[1]] == cols, "cabeçalho")
data = [[c.value for c in r] for r in ws2.iter_rows(min_row=2)]
check(len(data) == 25, f"{len(data)} linhas")
check([ascii_(r[0]) for r in data] == sorted(ascii_(r[0]) for r in data), "ordem alfabética")
por_mat = {r[1]: r for r in data}
for r in base:
    check(por_mat.get(r[1]) == r, f"{r[1]} difere do E5")
check(len({r[0] for r in data}) == 25 and len({r[4] for r in data}) == 25 and len({r[6] for r in data}) == 25,
      "nome/e-mail/telefone repetido")
for r in data:
    m, tel, lp = r[1], r[6], r[8]
    check(isinstance(m, str) and len(m) == 8 and m.isdigit(), f"matrícula {m}")
    check(isinstance(tel, str) and tel.isdigit() and len(tel) == 13, f"{m}: telefone")
    check(lp is None or (isinstance(lp, str) and lp <= "2026-10-15"), f"{m}: LP {lp!r}")
    check(isinstance(r[7], int) and 19 <= r[7] <= 52, f"{m}: idade")
conhecidos = {r[1] for r in base} | {"99000022"}           # S22 já tinha saído no E4
novos = [m for m in por_mat if m not in conhecidos]
print("Linhas:", len(data), "| novos:", sorted(novos), "| atualizados:", len(data) - len(novos),
      "| membros Yes:", sum(r[11] == "Yes" for r in data))
print("ERROS:" if erros else "OK: todas as validações passaram")
for e in erros: print(" -", e)
sys.exit(1 if erros else 0)
